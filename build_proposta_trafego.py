from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Proposta_Comercial_Trafego_Pago_Gatu_Digital.docx"
CAT = ROOT / "assets" / "gatu-digital-cat.png"

INK = "17131A"
PURPLE = "7E0078"
PURPLE_DARK = "52004F"
LILAC = "F5E6F4"
WASH = "FBF9FB"
MUTED = "6C6670"
LINE = "E9E5E8"
WHITE = "FFFFFF"

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(.72)
sec.bottom_margin = Inches(.68)
sec.left_margin = Inches(.82)
sec.right_margin = Inches(.82)
sec.header_distance = Inches(.3)
sec.footer_distance = Inches(.3)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc.get_or_add_tcPr()
    tcMar = tc.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tc.append(tcMar)
    for tag, val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+tag))
        if node is None: node=OxmlElement('w:'+tag); tcMar.append(node)
        node.set(qn('w:w'), str(val)); node.set(qn('w:type'),'dxa')

def set_cell_width(cell, dxa):
    tcPr = cell._tc.get_or_add_tcPr(); tcW = tcPr.find(qn('w:tcW'))
    if tcW is None: tcW=OxmlElement('w:tcW'); tcPr.append(tcW)
    tcW.set(qn('w:w'), str(dxa)); tcW.set(qn('w:type'),'dxa')

def table_geometry(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None: tblW=OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), str(sum(widths))); tblW.set(qn('w:type'),'dxa')
    ind = tblPr.find(qn('w:tblInd'))
    if ind is None: ind=OxmlElement('w:tblInd'); tblPr.append(ind)
    ind.set(qn('w:w'),'0'); ind.set(qn('w:type'),'dxa')
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for w in widths:
        col=OxmlElement('w:gridCol'); col.set(qn('w:w'),str(w)); grid.append(col)
    for row in table.rows:
        for c,w in zip(row.cells,widths): set_cell_width(c,w)

def no_borders(table):
    tblPr=table._tbl.tblPr; borders=tblPr.find(qn('w:tblBorders'))
    if borders is None: borders=OxmlElement('w:tblBorders'); tblPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'nil'); borders.append(e)

def font(run, size=10.5, color=INK, bold=False, italic=False, name='Arial'):
    run.font.name=name; run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),name)
    run.font.size=Pt(size); run.font.color.rgb=RGBColor.from_string(color)
    run.bold=bold; run.italic=italic
    return run

styles = doc.styles
normal=styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string(INK)
normal.paragraph_format.space_after=Pt(7); normal.paragraph_format.line_spacing=1.16
for sname,size,color,before,after in [('Title',34,INK,0,10),('Heading 1',23,INK,8,12),('Heading 2',14,PURPLE,10,5),('Heading 3',11,INK,6,3)]:
    s=styles[sname]; s.font.name='Arial'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color)
    s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True

def para(text='', size=10.5, color=INK, bold=False, italic=False, align=None, before=0, after=7, style=None):
    p=doc.add_paragraph(style=style); p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.16
    if align is not None: p.alignment=align
    font(p.add_run(text),size,color,bold,italic)
    return p

def heading(text, level=1):
    return doc.add_paragraph(text, style=f'Heading {level}')

def kicker(text):
    p=para(text.upper(),7.5,PURPLE,True,False,None,0,8)
    p.runs[0].font.letter_spacing if hasattr(p.runs[0].font,'letter_spacing') else None
    return p

def rule(color=LINE, size=10):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(8)
    pPr=p._p.get_or_add_pPr(); bdr=OxmlElement('w:pBdr'); bot=OxmlElement('w:bottom')
    bot.set(qn('w:val'),'single'); bot.set(qn('w:sz'),str(size)); bot.set(qn('w:color'),color); bdr.append(bot); pPr.append(bdr)

def pill(text, fill=LILAC, color=PURPLE):
    t=doc.add_table(rows=1, cols=1); table_geometry(t,[9360]); no_borders(t)
    c=t.cell(0,0); shade(c,fill); margins(c,150,220,150,220)
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); font(p.add_run(text),11,color,True)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def page_break(): doc.add_page_break()

def label_value(label, value):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3)
    font(p.add_run(label+' '),9,MUTED,True); font(p.add_run(value),9,INK,False)

def footer(section):
    h=section.header.paragraphs[0]; h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    font(h.add_run('gatu.'),8,INK,True); font(h.add_run('digital'),8,PURPLE,True)
    hPr=h._p.get_or_add_pPr(); hBdr=OxmlElement('w:pBdr'); hBot=OxmlElement('w:bottom')
    hBot.set(qn('w:val'),'single'); hBot.set(qn('w:sz'),'10'); hBot.set(qn('w:color'),PURPLE); hBdr.append(hBot); hPr.append(hBdr)
    f=section.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    font(f.add_run('gatu.digital  |  Proposta comercial de tráfego pago  |  '),8,MUTED)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)

footer(sec)

# CAPA
kicker('proposta comercial')
cover=doc.add_table(rows=1,cols=1); table_geometry(cover,[9360]); no_borders(cover)
cc=cover.cell(0,0); shade(cc,INK); margins(cc,430,420,430,420)
p=cc.paragraphs[0]; p.paragraph_format.space_after=Pt(48)
font(p.add_run('gatu.'),15,WHITE,True); font(p.add_run('digital'),15,'C946C1',True)
p=cc.add_paragraph(); p.paragraph_format.space_after=Pt(8); font(p.add_run('TRÁFEGO'),34,WHITE,True)
p=cc.add_paragraph(); p.paragraph_format.space_after=Pt(20); font(p.add_run('PAGO'),34,'C946C1',True)
p=cc.add_paragraph(); p.paragraph_format.space_after=Pt(0)
font(p.add_run('Estratégia e dados conectados para transformar mídia em oportunidades reais de negócio.'),14,'E9E5E8')
doc.add_paragraph().paragraph_format.space_after=Pt(6)
t=doc.add_table(rows=1,cols=1); table_geometry(t,[9360]); no_borders(t)
c=t.cell(0,0); margins(c,120,0,120,0)
p=c.paragraphs[0]; font(p.add_run('Preparado para\n'),9,MUTED,True); font(p.add_run('[NOME DO CLIENTE / EMPRESA]'),13,INK,True)
p=c.add_paragraph(); font(p.add_run('Objetivo\n'),9,MUTED,True); font(p.add_run('Atrair, medir e converter demanda com Meta Ads + Google Ads'),10.5,INK)
para('Estratégia. Aquisição. Conversão. Operação.',10,PURPLE,True,False,None,26,3)
para('Validade da proposta: 10 dias  |  Data: [DD/MM/AAAA]',9,MUTED)

# P2
page_break(); rule(PURPLE,8); heading('Anunciar é fácil. Construir uma máquina de aquisição é outra história.',1)
para('Campanhas sem estrutura, rastreamento e rotina de otimização podem até gerar cliques — mas não dão clareza sobre o que realmente trouxe conversas, oportunidades e vendas. A proposta da gatu.digital é cuidar da jornada completa: da fundação técnica à decisão orientada por dados.',12,MUTED,False,False,None,0,14)
pill('O objetivo não é “gastar em anúncios”. É comprar aprendizado, acelerar oportunidades e tornar cada decisão mais inteligente.')
heading('O valor que você passa a enxergar',2)
items=[('CLAREZA','Saber quais campanhas, anúncios, públicos e palavras-chave estão produzindo oportunidades reais.'),('CONTROLE','Ter contas, verificações, acessos e mensuração organizados para reduzir riscos e desperdícios.'),('VELOCIDADE','Aprender mais rápido com testes, leitura de performance e otimizações contínuas.'),('CONSISTÊNCIA','Manter criação, mídia, rastreamento e comunicação trabalhando como uma única operação.')]
t=doc.add_table(rows=2,cols=2); table_geometry(t,[4680,4680]); no_borders(t)
for cell,(lab,txt) in zip([c for row in t.rows for c in row.cells],items):
    shade(cell,WASH); margins(cell,210,220,210,220); p=cell.paragraphs[0]; font(p.add_run(lab),8,PURPLE,True); p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(0); font(p.add_run(txt),10,INK)
heading('Resultado esperado',2)
para('Uma estrutura de mídia paga capaz de atrair o público certo, medir o caminho até a conversão e melhorar continuamente a relação entre investimento, volume e qualidade das oportunidades.',11,INK,True)

# P3
page_break(); kicker('metodologia gatu'); heading('Um caminho claro para crescer.',1)
para('A gestão é organizada em quatro movimentos conectados. Cada etapa resolve um risco que, quando ignorado, compromete as seguintes.',11,MUTED,False,False,None,0,14)
steps=[('G','GUIAR','Diagnóstico, metas, oferta, público e indicadores. Sem direção, a campanha otimiza o que é fácil medir — não necessariamente o que faz o negócio crescer.'),('A','ATRAIR','Meta Ads, Google Ads, palavras-chave, criativos e copy. Aqui transformamos estratégia em demanda qualificada.'),('T','TRANSFORMAR','Rastreamento do anúncio ao WhatsApp ou página, leitura do funil e otimização para conversão. Clique só tem valor quando avança a jornada.'),('U','UNIFICAR','Dados, feedback semanal e próximos testes em uma rotina de melhoria. O aprendizado volta para as campanhas e fortalece o próximo ciclo.')]
for letter,title,txt in steps:
    t=doc.add_table(rows=1,cols=2); table_geometry(t,[900,8460]); no_borders(t)
    c1,c2=t.rows[0].cells; shade(c1,PURPLE); margins(c1,200,100,200,100); c1.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=c1.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run(letter),20,WHITE,True)
    shade(c2,WASH); margins(c2,180,220,180,220); p=c2.paragraphs[0]; font(p.add_run(title),10,PURPLE,True); p=c2.add_paragraph(); p.paragraph_format.space_after=Pt(0); font(p.add_run(txt),10,INK)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
pill('A força está na conexão: estratégia orienta a mídia; mídia gera dados; dados revelam gargalos; otimização transforma aprendizado em resultado.',INK,WHITE)

# P4
page_break(); rule(PURPLE,8); heading('Dez frentes. Uma única responsabilidade: fazer a operação performar.',1)
scope=[
('01','Setup de Meta Ads e Google Ads','Organização das contas, ativos, acessos, formas de pagamento e estrutura necessária para operar com segurança.','Evita bloqueios, perda de histórico e campanhas construídas sobre uma base frágil.'),
('02','Verificação da empresa','Apoio no processo de verificação e validação da empresa nas plataformas.','Aumenta confiabilidade, segurança e acesso a recursos importantes.'),
('03','Planejamento e criação de campanhas','Definição de objetivos, públicos, ofertas, canais, orçamento, estrutura e testes.','Transforma intenção em um plano executável, com hipóteses claras.'),
('04','Análise de performance','Leitura de indicadores por campanha, anúncio, público, palavra-chave e etapa do funil.','Mostra o que escalar, corrigir ou interromper antes que o desperdício cresça.'),
('05','Arte, copy e legenda dos anúncios','Criação das peças e mensagens necessárias para comunicar valor e gerar ação.','Mídia compra atenção; a criatividade converte atenção em interesse.'),
]
for num,title,doing,value in scope:
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(2); p.paragraph_format.keep_with_next=True
    font(p.add_run(num+'  '),9,PURPLE,True); font(p.add_run(title),11.5,INK,True)
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.38); p.paragraph_format.space_after=Pt(2); font(p.add_run(doing),9.7,INK)
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.38); p.paragraph_format.space_after=Pt(5); font(p.add_run('POR QUE É INDISPENSÁVEL  '),8,PURPLE,True); font(p.add_run(value),9.3,MUTED)

# P5
page_break(); rule(PURPLE,8); heading('Inteligência para evoluir toda semana.',1)
scope2=[
('06','Pesquisa de palavras-chave e concorrentes','Mapeamento de intenção de busca, termos prioritários, negativas, oportunidades e movimentos do mercado.','Coloca a marca diante de quem já está procurando e evita pagar por buscas sem potencial.'),
('07','Acompanhamento e otimização','Ajustes de orçamento, lances, segmentação, posicionamentos, criativos e termos de busca.','Campanha não é peça pronta: performance nasce da combinação entre dados, decisão e recorrência.'),
('08','Feedback semanal','Resumo objetivo do que aconteceu, aprendizados, ações realizadas e próximos testes.','Mantém transparência, alinha expectativas e acelera decisões do cliente.'),
('09','Rastreamento de campanhas','Configuração de UTMs, eventos, pixels/tags e identificação da origem das conversões.','Sem rastreamento, a plataforma aprende com sinais incompletos e o negócio investe no escuro.'),
('10','Funil de aquisição','Visão da jornada entre impressão, clique, visita, conversa, lead e venda — conforme os dados disponíveis.','Conecta marketing ao resultado comercial e revela onde as oportunidades estão sendo perdidas.')]
for num,title,doing,value in scope2:
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(2); p.paragraph_format.keep_with_next=True
    font(p.add_run(num+'  '),9,PURPLE,True); font(p.add_run(title),11.5,INK,True)
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.38); p.paragraph_format.space_after=Pt(2); font(p.add_run(doing),9.7,INK)
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.38); p.paragraph_format.space_after=Pt(5); font(p.add_run('POR QUE É INDISPENSÁVEL  '),8,PURPLE,True); font(p.add_run(value),9.3,MUTED)
pill('Diferencial gatu.digital: rastreamento das campanhas do WhatsApp incluído na gestão, ampliando a inteligência sobre a origem das conversas.')

# P6
page_break(); rule(PURPLE,8); heading('Do anúncio à oportunidade: o caminho precisa ser visível.',1)
funnel=[('1','ATENÇÃO','Impressões e alcance'),('2','INTERESSE','Cliques e engajamento'),('3','INTENÇÃO','Visitas, buscas e ações'),('4','CONVERSA','WhatsApp, formulário ou ligação'),('5','RESULTADO','Lead qualificado e venda*')]
for n,lab,metric in funnel:
    t=doc.add_table(rows=1,cols=3); table_geometry(t,[600,2500,6260]); no_borders(t)
    c1,c2,c3=t.rows[0].cells; shade(c1,PURPLE); shade(c2,LILAC); shade(c3,WASH)
    for c in (c1,c2,c3): margins(c,120,150,120,150)
    p=c1.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run(n),10,WHITE,True)
    font(c2.paragraphs[0].add_run(lab),9,PURPLE,True); font(c3.paragraphs[0].add_run(metric),9.5,INK)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
para('*A leitura de vendas depende do retorno comercial e/ou integração disponível.',8,MUTED,True,False,None,3,10)
heading('Cadência de trabalho',2)
cad=[('INÍCIO','Diagnóstico, acessos, verificações, mensuração, pesquisa e plano de campanha.'),('SEMANAL','Monitoramento, otimizações, análise e feedback com aprendizados e próximos passos.'),('MENSAL','Leitura consolidada, evolução dos indicadores e priorização do novo ciclo de testes.')]
t=doc.add_table(rows=1,cols=3); table_geometry(t,[3120,3120,3120]); no_borders(t)
for c,(lab,txt) in zip(t.rows[0].cells,cad):
    shade(c,WASH); margins(c,180,180,180,180); font(c.paragraphs[0].add_run(lab),8,PURPLE,True); p=c.add_paragraph(); p.paragraph_format.space_after=Pt(0); font(p.add_run(txt),9.3,INK)
heading('Indicadores acompanhados',2)
para('Investimento, alcance, CTR, CPC, CPM, conversões, custo por conversão/lead, termos de busca, taxa de avanço no funil e qualidade das oportunidades — escolhidos conforme objetivo e disponibilidade de dados.',10,MUTED)

# P7
page_break(); kicker('investimento'); heading('Uma gestão completa para fazer cada real trabalhar com mais inteligência.',1)
t=doc.add_table(rows=1,cols=2); table_geometry(t,[5700,3660]); no_borders(t)
c1,c2=t.rows[0].cells; shade(c1,INK); shade(c2,PURPLE)
for c in (c1,c2): margins(c,260,260,260,260)
font(c1.paragraphs[0].add_run('GESTÃO MENSAL'),9,WHITE,True); p=c1.add_paragraph(); font(p.add_run('Meta Ads + Google Ads'),18,WHITE,True); p=c1.add_paragraph(); p.paragraph_format.space_after=Pt(0); font(p.add_run('Estratégia, criação, operação, análise, otimização, feedback semanal e rastreamento.'),10,'D6CAD1')
font(c2.paragraphs[0].add_run('INVESTIMENTO'),9,WHITE,True); p=c2.add_paragraph(); font(p.add_run('A partir de'),9,WHITE,True); p=c2.add_paragraph(); font(p.add_run('R$ 1.200'),22,WHITE,True); p=c2.add_paragraph(); p.paragraph_format.space_after=Pt(0); font(p.add_run('/ mês'),9,WHITE,True)
heading('Condições comerciais',2)
conditions=[
('Verba de mídia','Paga diretamente pelo cliente às plataformas; não está incluída na gestão.'),
('Prazo recomendado','Ciclo inicial mínimo sugerido de 3 meses para estruturar, aprender e otimizar com consistência.'),
('Escopo criativo','Peças e variações necessárias ao plano de mídia, conforme pauta e materiais aprovados.'),
('Acessos e aprovações','O cliente disponibiliza acessos, informações, materiais e retornos necessários ao andamento.'),
('Resultados','Não há promessa de resultado fixo: performance depende de oferta, mercado, verba, atendimento, histórico e velocidade de aprovação.')]
for lab,val in conditions: label_value(lab+' —',val)
pill('Próximo passo: aprovação da proposta → briefing e acessos → setup e rastreamento → campanhas no ar → otimização contínua.',LILAC,PURPLE)
heading('Vamos construir uma operação que aprende e melhora a cada campanha?',2)
para('A gatu.digital conecta estratégia, performance e tecnologia para que o seu negócio seja encontrado, escolhido e transformado em oportunidades reais.',11,INK,True)
para('WhatsApp: (31) 97259-0277  |  gatu.digital',10,PURPLE,True,False,None,5,2)
para('ACEITE:  ____________________________________    DATA: ____ / ____ / ______',9,MUTED,True,False,None,20,0)

# Document properties
doc.core_properties.title='Proposta Comercial de Tráfego Pago — gatu.digital'
doc.core_properties.subject='Gestão de Meta Ads e Google Ads com rastreamento e otimização'
doc.core_properties.author='gatu.digital'
doc.core_properties.keywords='tráfego pago, Meta Ads, Google Ads, rastreamento, performance'

# Pre-flight editorial: the selected design skill bans visible em dashes.
def replace_em_dashes(container):
    for p in container.paragraphs:
        for r in p.runs:
            if '—' in r.text:
                r.text = r.text.replace('—', ':')
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                replace_em_dashes(cell)

replace_em_dashes(doc)
doc.save(OUT)
print(OUT)
